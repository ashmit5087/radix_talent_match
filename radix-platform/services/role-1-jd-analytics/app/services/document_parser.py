"""Turn an uploaded PDF/DOCX (as bytes) into plain text.

Safety notes:
  * Nothing is written to disk: files are parsed from memory (BytesIO), so there are no temp
    files to clean up and no original upload is ever stored.
  * Only text is read. PDFs are read with pypdf and DOCX with python-docx; neither executes
    macros, scripts or embedded content.
"""
import io
import logging
import os
import re
import zipfile
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader

from app.config import Settings
from app.errors import AppError

logger = logging.getLogger("radix.parser")

ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx"}
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024  # zip-bomb guard


@dataclass
class ParsedDocument:
    filename: str
    kind: str  # "pdf" | "docx"
    text: str


def sanitize_filename(name: str | None) -> str:
    """Keep only the bare file name (no directories), so no local path can leak."""
    name = (name or "").replace("\\", "/").split("/")[-1].strip()
    name = re.sub(r"[\x00-\x1f]", "", name)
    return name[:255] or "unnamed"


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _detect_kind(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    kind = ALLOWED_EXTENSIONS.get(ext)
    if kind is None:
        raise AppError(415, "UNSUPPORTED_FILE_TYPE", "Upload a PDF or DOCX file.")
    if not content:
        raise AppError(400, "EMPTY_FILE", "The uploaded file is empty.")
    # Check the file's real signature, not just its extension.
    if kind == "pdf" and not content.lstrip()[:5].startswith(b"%PDF-"):
        raise AppError(400, "CORRUPT_DOCUMENT", "The file is not a valid PDF.")
    if kind == "docx" and not content.startswith(b"PK"):
        raise AppError(400, "CORRUPT_DOCUMENT", "The file is not a valid DOCX.")
    return kind


def _extract_pdf(content: bytes, settings: Settings) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")  # many PDFs are "encrypted" with an empty password
            except Exception:
                ok = 0
            if not ok:
                raise AppError(400, "ENCRYPTED_DOCUMENT", "The PDF is password-protected and cannot be read.")
        if len(reader.pages) > settings.max_pdf_pages:
            raise AppError(413, "TOO_MANY_PAGES", f"The PDF has more than {settings.max_pdf_pages} pages.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except AppError:
        raise
    except Exception:
        raise AppError(400, "CORRUPT_DOCUMENT", "The PDF is corrupt or unreadable.") from None


def _extract_docx(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise AppError(413, "DOCUMENT_TOO_LARGE", "The DOCX expands to an unreasonable size.")
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) not in seen:  # merged cells repeat; read each once
                        seen.add(id(cell._tc))
                        parts.append(cell.text)
        return "\n".join(parts)
    except AppError:
        raise
    except Exception:
        raise AppError(400, "CORRUPT_DOCUMENT", "The DOCX is corrupt or unreadable.") from None


def parse_document(filename: str, content: bytes, settings: Settings) -> ParsedDocument:
    """Validate and extract text. Raises AppError with a clear message on any problem."""
    filename = sanitize_filename(filename)
    if len(content) > settings.max_upload_bytes:
        raise AppError(413, "FILE_TOO_LARGE", f"Upload exceeds the {settings.max_upload_size_mb:g} MB limit.")
    kind = _detect_kind(filename, content)
    raw = _extract_pdf(content, settings) if kind == "pdf" else _extract_docx(content)
    text = _clean_text(raw)
    if not text:
        raise AppError(
            400, "NO_TEXT_FOUND",
            "No text could be extracted. Scanned/image-only documents are not supported (no OCR).",
        )
    if len(text) > settings.max_extracted_text_chars:
        raise AppError(
            413, "EXTRACTED_TEXT_TOO_LARGE",
            f"Extracted text exceeds the {settings.max_extracted_text_chars}-character limit.",
        )
    logger.info("Parsed %s (%s): %d chars", filename, kind, len(text))  # never log the text itself
    return ParsedDocument(filename=filename, kind=kind, text=text)
