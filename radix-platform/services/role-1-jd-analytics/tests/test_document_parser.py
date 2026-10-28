import io

import pytest
from pypdf import PdfReader, PdfWriter

from app.errors import AppError
from app.services.document_parser import parse_document, sanitize_filename
from tests.conftest import ML_JD, make_blank_pdf, make_docx, make_pdf


def test_pdf_extraction(settings):
    doc = parse_document("jd.pdf", make_pdf(ML_JD), settings)
    assert doc.kind == "pdf" and "machine learning" in doc.text and "Python" in doc.text


def test_docx_extraction_includes_tables(settings):
    from docx import Document
    d = Document()
    d.add_paragraph("Responsibilities: write Python")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Skill"
    t.cell(0, 1).text = "Kubernetes"
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_document("jd.docx", buf.getvalue(), settings)
    assert "write Python" in doc.text and "Kubernetes" in doc.text


@pytest.mark.parametrize("name,content,status,code", [
    ("jd.txt", b"hello", 415, "UNSUPPORTED_FILE_TYPE"),
    ("jd.pdf", b"", 400, "EMPTY_FILE"),
    ("jd.pdf", b"not really a pdf", 400, "CORRUPT_DOCUMENT"),
    ("jd.docx", b"not a zip", 400, "CORRUPT_DOCUMENT"),
    ("jd.docx", b"PK\x03\x04garbage-garbage", 400, "CORRUPT_DOCUMENT"),
    ("jd.pdf", b"%PDF-1.4 truncated garbage", 400, "CORRUPT_DOCUMENT"),
])
def test_invalid_files(settings, name, content, status, code):
    with pytest.raises(AppError) as e:
        parse_document(name, content, settings)
    assert (e.value.status_code, e.value.code) == (status, code)


def test_textless_pdf(settings):
    with pytest.raises(AppError) as e:
        parse_document("scan.pdf", make_blank_pdf(), settings)
    assert e.value.code == "NO_TEXT_FOUND"


def test_encrypted_pdf(settings):
    w = PdfWriter()
    for p in PdfReader(io.BytesIO(make_pdf(ML_JD))).pages:
        w.add_page(p)
    w.encrypt("s3cret", algorithm="RC4-128")
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(AppError) as e:
        parse_document("locked.pdf", buf.getvalue(), settings)
    assert e.value.code == "ENCRYPTED_DOCUMENT"


def test_oversized_upload(settings):
    with pytest.raises(AppError) as e:
        parse_document("big.pdf", b"%PDF-" + b"0" * (2 * 1024 * 1024), settings)
    assert e.value.status_code == 413 and e.value.code == "FILE_TOO_LARGE"


def test_extracted_text_limit(settings):
    s = settings.model_copy(update={"max_extracted_text_chars": 1000})
    with pytest.raises(AppError) as e:
        parse_document("long.docx", make_docx("word " * 500), s)
    assert e.value.status_code == 413 and e.value.code == "EXTRACTED_TEXT_TOO_LARGE"


def test_filename_never_contains_path():
    assert sanitize_filename("C:\\Users\\me\\docs\\jd.pdf") == "jd.pdf"
    assert sanitize_filename("/tmp/x/../jd.pdf") == "jd.pdf"
